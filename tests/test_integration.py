"""End-to-end integration test running the complete ReportSmith pipeline."""

from pathlib import Path
import docx
import pytest

from reportsmith.brief import load_brief
from reportsmith.check import validate_report
from reportsmith.config import load_config
from reportsmith.diagrams import generate_all_diagrams
from reportsmith.docx_build import DocxReportBuilder
from reportsmith.facts import derive_facts
from reportsmith.generate import SectionGenerator
from reportsmith.llm import StubLLM
from reportsmith.outline import parse_format_file
from reportsmith.pagenums import find_libreoffice


def test_full_pipeline_stub_llm(tmp_path):
    brief_file = Path("examples/brief.yaml")
    format_file = Path("examples/format.txt")
    sample_file = Path("examples/sample.docx")

    assert brief_file.is_file()
    assert format_file.is_file()
    assert sample_file.is_file()

    # 1. Load brief and facts
    brief = load_brief(brief_file)
    facts = derive_facts(brief, base_dir=brief_file.parent)

    # 2. Parse outline
    outline = parse_format_file(format_file)
    assert len(outline.items) > 20

    # 3. Generate diagrams
    diag_dir = tmp_path / "diagrams"
    diag_dir.mkdir(parents=True, exist_ok=True)
    diagrams = generate_all_diagrams(brief, facts, out_dir=diag_dir)
    assert Path(diagrams["block_diagram"]).is_file()
    assert Path(diagrams["connection_diagram"]).is_file()
    assert Path(diagrams["flowchart"]).is_file()

    # 4. Generate sections with StubLLM
    cfg = load_config()
    cfg.paths.cache_dir = str(tmp_path / "cache")
    stub = StubLLM(seed=42)
    generator = SectionGenerator(stub, config=cfg, cache_dir=tmp_path / "cache")
    sections = generator.generate_all(outline, brief, facts, resume=False)
    assert len(sections) > 10

    # 5. Build docx
    out_docx = tmp_path / "integration_report.docx"
    builder = DocxReportBuilder(
        brief=brief,
        facts=facts,
        outline=outline,
        sample_path=sample_file,
    )
    doc = builder.build(sections, out_path=out_docx)
    assert out_docx.is_file()

    # Verify docx paragraphs & tables
    loaded_doc = docx.Document(str(out_docx))
    assert len(loaded_doc.paragraphs) > 50
    assert len(loaded_doc.tables) >= 5

    # 6. Validate with check.py
    check_txt = tmp_path / "check_summary.txt"
    is_valid = validate_report(out_docx, brief, facts, out_report_path=check_txt)
    assert is_valid is True

    # 7. Test LibreOffice if installed
    soffice_path = find_libreoffice()
    if soffice_path:
        import subprocess
        from pypdf import PdfReader
        pdf_out = tmp_path / "integration_report.pdf"
        cmd = [str(soffice_path), "--headless", "--convert-to", "pdf", str(out_docx), "--outdir", str(tmp_path)]
        res = subprocess.run(cmd, capture_output=True, timeout=120)
        assert res.returncode == 0
        assert pdf_out.is_file()
        reader = PdfReader(str(pdf_out))
        assert len(reader.pages) > 0
