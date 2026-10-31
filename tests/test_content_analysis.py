"""Test suite for Content folder analysis and COA_report.docx styling pipeline."""

from pathlib import Path
import docx
import pytest

from reportsmith.check import validate_report
from reportsmith.config import load_config
from reportsmith.content_analyzer import analyze_content_folder
from reportsmith.docx_build import DocxReportBuilder
from reportsmith.generate import SectionGenerator
from reportsmith.llm import StubLLM


def test_content_folder_analysis():
    brief, facts, outline = analyze_content_folder("Content", "COA_report.docx")

    assert "Water" in brief.project.title or "Plant" in brief.project.title
    assert "CST504" in brief.project.course_code
    assert brief.implementation.board.mcu_part_number == "STM32F103C8T6"
    assert len(brief.people.team) == 4
    assert len(outline.items) >= 40

    # Verify pins
    pin_map = {p["pin"]: p for p in facts.pin_table}
    assert "PC13" in pin_map
    assert any(p in pin_map for p in ("PA0", "PA1", "PB0", "PB1"))

    # Verify state ranges
    assert len(facts.state_ranges) >= 3


def test_end_to_end_content_build_and_validation(tmp_path):
    brief, facts, outline = analyze_content_folder("Content", "COA_report.docx")
    config = load_config()

    # Generate sections using deterministic stub LLM
    generator = SectionGenerator(StubLLM(), config=config, cache_dir=tmp_path / "cache")
    gen_data = generator.generate_all(outline, brief, facts)

    out_docx = tmp_path / "test_report.docx"
    builder = DocxReportBuilder(brief, facts, outline, sample_path="COA_report.docx")
    builder.build(gen_data, out_docx)

    assert out_docx.is_file()
    assert out_docx.stat().st_size > 50000  # At least 50 KB

    # Inspect generated docx
    doc = docx.Document(str(out_docx))
    assert len(doc.sections) == 3

    # Check Title Page
    title_p = doc.paragraphs[0]
    assert "WATER" in title_p.text.upper() or "STM32" in title_p.text.upper()
    assert title_p.runs[0].font.size.pt == 26.0

    # Check Team Table (Table 0)
    team_tbl = doc.tables[0]
    assert len(team_tbl.rows) == 4
    assert len(team_tbl.rows[0].cells[0].text.strip()) > 0

    # Check Certificate
    cert_headings = [p.text for p in doc.paragraphs if "BONAFIDE CERTIFICATE" in p.text]
    assert len(cert_headings) >= 1

    # Check Contents Banner Table
    banner_tables = [t for t in doc.tables if len(t.rows) == 1 and len(t.columns) == 1 and "CONTENTS" in t.rows[0].cells[0].text]
    assert len(banner_tables) >= 1

    # Check Verbatim Source Code Appendix
    code_blocks = [
        c.text
        for t in doc.tables
        for r in t.rows
        for c in r.cells
        if len(c.text) > 50 and ("setup" in c.text or "loop" in c.text or "void" in c.text or "pinMode" in c.text)
    ]
    assert len(code_blocks) >= 1

    # Validate with grounding checker
    check_report = tmp_path / "check.txt"
    is_valid = validate_report(out_docx, brief, facts, out_report_path=check_report)
    assert is_valid is True
    assert "Status: PASSED" in check_report.read_text(encoding="utf-8")
