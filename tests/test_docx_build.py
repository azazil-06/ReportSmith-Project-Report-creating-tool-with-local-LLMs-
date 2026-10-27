"""Unit tests for Word (.docx) report builder."""

from pathlib import Path
import docx
import pytest

from reportsmith.brief import load_brief
from reportsmith.docx_build import DocxReportBuilder
from reportsmith.facts import derive_facts
from reportsmith.generate import SectionGenerator
from reportsmith.llm import StubLLM
from reportsmith.outline import parse_format_file


def test_docx_build_structure_and_sections(tmp_path):
    brief = load_brief("examples/brief.yaml")
    facts = derive_facts(brief, base_dir=Path("."))
    outline = parse_format_file("examples/format.txt")

    from reportsmith.config import load_config
    cfg = load_config()
    cfg.paths.cache_dir = str(tmp_path / "cache")
    stub = StubLLM()
    gen = SectionGenerator(stub, config=cfg, cache_dir=tmp_path / "cache")

    generated_sections = gen.generate_all(
        outline=outline,
        brief=brief,
        facts=facts,
        resume=False,
    )

    out_file = tmp_path / "test_report.docx"
    builder = DocxReportBuilder(
        brief=brief,
        facts=facts,
        outline=outline,
        sample_path="examples/sample.docx",
    )
    doc = builder.build(generated_sections, out_path=out_file)

    assert out_file.is_file()
    loaded_doc = docx.Document(str(out_file))

    # Test 3 sections: Title Page, Prelims, Body
    assert len(loaded_doc.sections) == 3

    # Check prelim section roman page numbering
    s2 = loaded_doc.sections[1]
    pgNumType2 = s2._sectPr.xpath(".//w:pgNumType")
    assert len(pgNumType2) > 0
    assert pgNumType2[0].attrib.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fmt") == "lowerRoman"

    # Check body section decimal page numbering restarting at 1
    s3 = loaded_doc.sections[2]
    pgNumType3 = s3._sectPr.xpath(".//w:pgNumType")
    assert len(pgNumType3) > 0
    assert pgNumType3[0].attrib.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fmt") == "decimal"
    assert pgNumType3[0].attrib.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}start") == "1"


def test_code_appendix_verbatim_check(tmp_path):
    brief = load_brief("examples/brief.yaml")
    facts = derive_facts(brief, base_dir=Path("."))
    outline = parse_format_file("examples/format.txt")

    from reportsmith.config import load_config
    cfg = load_config()
    cfg.paths.cache_dir = str(tmp_path / "cache")
    stub = StubLLM()
    gen = SectionGenerator(stub, config=cfg, cache_dir=tmp_path / "cache")

    generated = gen.generate_all(outline=outline, brief=brief, facts=facts, resume=False)
    out_file = tmp_path / "appendix_test.docx"
    builder = DocxReportBuilder(brief=brief, facts=facts, outline=outline)
    builder.build(generated, out_path=out_file)

    loaded = docx.Document(str(out_file))

    # Read original source code
    original_code = Path("examples/plant_monitor.ino").read_text(encoding="utf-8")

    # Check that code block in docx contains key lines verbatim
    all_text = "\n".join(p.text for p in loaded.paragraphs)
    for table in loaded.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text += "\n" + cell.text

    assert "#define AIR_VALUE 3700" in all_text
    assert "#define WATER_VALUE 1300" in all_text
    assert "readSoilMoistureAverage()" in all_text
    assert "calculateMoisturePercent(int rawAdc)" in all_text
    assert "updateDisplay(rawAdc, percent, state)" in all_text


def test_tables_formatting_and_cant_split(tmp_path):
    brief = load_brief("examples/brief.yaml")
    facts = derive_facts(brief, base_dir=Path("."))
    outline = parse_format_file("examples/format.txt")

    from reportsmith.config import load_config
    cfg = load_config()
    cfg.paths.cache_dir = str(tmp_path / "cache")
    stub = StubLLM()
    gen = SectionGenerator(stub, config=cfg, cache_dir=tmp_path / "cache")

    generated = gen.generate_all(outline=outline, brief=brief, facts=facts, resume=False)
    out_file = tmp_path / "tables_test.docx"
    builder = DocxReportBuilder(brief=brief, facts=facts, outline=outline)
    builder.build(generated, out_path=out_file)

    loaded = docx.Document(str(out_file))

    # Check tables have cantSplit on trPr
    found_cant_split = False
    for table in loaded.tables:
        for row in table.rows:
            cant_splits = row._tr.xpath(".//w:cantSplit")
            if cant_splits:
                found_cant_split = True
                break
        if found_cant_split:
            break
    assert found_cant_split, "Table rows should have w:cantSplit to avoid page splits"
