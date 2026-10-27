"""Unit tests for validator (check.py) and planted error detection."""

from pathlib import Path
import docx
import pytest

from reportsmith.brief import load_brief
from reportsmith.check import validate_report, ReportValidator
from reportsmith.docx_build import DocxReportBuilder
from reportsmith.facts import derive_facts
from reportsmith.generate import SectionGenerator
from reportsmith.llm import StubLLM
from reportsmith.outline import parse_format_file


from reportsmith.config import load_config


def _get_or_create_report(out_path: Path):
    if not out_path.is_file():
        brief = load_brief("examples/brief.yaml")
        facts = derive_facts(brief, base_dir=Path("."))
        outline = parse_format_file("examples/format.txt")
        config = load_config()
        gen = SectionGenerator(StubLLM(), config=config)
        gen_data = gen.generate_all(outline, brief, facts)
        builder = DocxReportBuilder(brief, facts, outline, sample_path="examples/sample.docx")
        builder.build(gen_data, out_path)
    return out_path


def test_validator_passes_on_valid_report(tmp_path):
    brief = load_brief("examples/brief.yaml")
    facts = derive_facts(brief, base_dir=Path("."))
    outline = parse_format_file("examples/format.txt")

    report_file = tmp_path / "valid_report.docx"
    _get_or_create_report(report_file)
    assert report_file.is_file(), "report must exist"

    check_txt = tmp_path / "check_report.txt"
    is_valid = validate_report(report_file, brief, facts, out_report_path=check_txt)
    assert is_valid is True
    assert "Status: PASSED" in check_txt.read_text(encoding="utf-8")


def test_validator_catches_planted_wrong_number(tmp_path):
    brief = load_brief("examples/brief.yaml")
    facts = derive_facts(brief, base_dir=Path("."))
    outline = parse_format_file("examples/format.txt")

    report_file = tmp_path / "base_report.docx"
    _get_or_create_report(report_file)

    # Copy report and plant an ungrounded number in prose
    doc = docx.Document(str(report_file))
    p = doc.add_paragraph()
    p.add_run("The microcontroller was overclocked to 987654 MHz during physical testing.")

    tampered_docx = tmp_path / "tampered.docx"
    doc.save(str(tampered_docx))

    check_txt = tmp_path / "check_tampered.txt"
    is_valid = validate_report(tampered_docx, brief, facts, out_report_path=check_txt)
    assert is_valid is False

    report_content = check_txt.read_text(encoding="utf-8")
    assert "987654" in report_content
    assert "Ungrounded number" in report_content


def test_validator_catches_banned_ai_phrases(tmp_path):
    brief = load_brief("examples/brief.yaml")
    facts = derive_facts(brief, base_dir=Path("."))

    report_file = tmp_path / "base_report2.docx"
    _get_or_create_report(report_file)

    doc = docx.Document(str(report_file))
    p = doc.add_paragraph()
    p.add_run("As an AI language model, I recommend calibrating the sensor further.")

    tampered_docx = tmp_path / "banned_ai.docx"
    doc.save(str(tampered_docx))

    check_txt = tmp_path / "check_banned.txt"
    is_valid = validate_report(tampered_docx, brief, facts, out_report_path=check_txt)
    assert is_valid is False

    report_content = check_txt.read_text(encoding="utf-8")
    assert "as an ai" in report_content.lower() or "language model" in report_content.lower()
